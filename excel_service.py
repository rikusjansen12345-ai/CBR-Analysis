from pathlib import Path
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule

from config import (
    REPORT_OUTPUT_FOLDER,
    REPORT_FILENAME,
)


def create_workbook(records):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "CBR Analysis"

    headers = [
        "Name",
        "Customer Name",
        "Sales Stage",
        "Rep",
        "Date Created",
        "Expected Close Date",
        "TD less CD",
        "EID less CD",
        "Forecast Value",
        "Deposit Paid",
        "Balance Due",
    ]

    # -------------------------
    # Header row
    # -------------------------
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=column)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            start_color="0F4C81",
            end_color="0F4C81",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    worksheet.row_dimensions[1].height = 24

    # -------------------------
    # Data rows
    # -------------------------
    for row_number, fields in enumerate(records, start=2):

        worksheet[f"A{row_number}"] = fields.get("name", "")
        worksheet[f"B{row_number}"] = fields.get("customer_name", "")
        worksheet[f"C{row_number}"] = fields.get("sales_stage", "")
        worksheet[f"D{row_number}"] = fields.get("rep", "")

        # Date Created
        created = fields.get("date_created")

        if created:
            if isinstance(created, str):
                worksheet[f"E{row_number}"] = (
                    datetime.fromisoformat(created.replace("Z", "+00:00"))
                    .replace(tzinfo=None)
                )
            else:
                worksheet[f"E{row_number}"] = created
        else:
            worksheet[f"E{row_number}"] = None

        # Expected Close Date
        expected = fields.get("expected_close_date")

        if expected:
            if isinstance(expected, str):
                worksheet[f"F{row_number}"] = (
                    datetime.fromisoformat(expected.replace("Z", "+00:00"))
                    .replace(tzinfo=None)
                )
            else:
                worksheet[f"F{row_number}"] = expected
        else:
            worksheet[f"F{row_number}"] = None

        # TD less CD
        worksheet[f"G{row_number}"] = (
            f'=IF(C{row_number}="Approved Not Here","N/A",INT(TODAY()-E{row_number}))'
        )

        # EID less CD
        worksheet[f"H{row_number}"] = (
            f'=IF(F{row_number}="","N/A",INT(F{row_number}-E{row_number}))'
        )

        # Forecast Value
        forecast = fields.get("forecast_value")

        if forecast is None or forecast == "":
            worksheet[f"I{row_number}"] = "N/A"
        else:
            worksheet[f"I{row_number}"] = forecast

        # Deposit Paid
        deposit = fields.get("deposit_paid")

        if deposit is None or deposit == "":
            worksheet[f"J{row_number}"] = "N/A"
        else:
            worksheet[f"J{row_number}"] = deposit

        # Balance Due
        if forecast is None or forecast == "":
            worksheet[f"K{row_number}"] = "N/A"
        elif deposit is None or deposit == "":
            worksheet[f"K{row_number}"] = "N/A"
        else:
            worksheet[f"K{row_number}"] = f"=I{row_number}-J{row_number}"

    # -------------------------
    # Totals Row
    # -------------------------
    total_row = worksheet.max_row + 1

    worksheet[f"H{total_row}"] = "TOTALS"
    worksheet[f"H{total_row}"].font = Font(bold=True)

    worksheet[f"I{total_row}"] = f"=SUM(I2:I{total_row-1})"
    worksheet[f"J{total_row}"] = f"=SUM(J2:J{total_row-1})"
    worksheet[f"K{total_row}"] = f"=SUM(K2:K{total_row-1})"

    for cell in [f"I{total_row}", f"J{total_row}", f"K{total_row}"]:
        worksheet[cell].font = Font(bold=True)
        worksheet[cell].number_format = "R #,##0.00"
    # -------------------------
    # Create Excel Table
    # -------------------------
    table = Table(
        displayName="CBRAnalysis",
        ref=f"A1:K{total_row-1}"
    )

    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    table.tableStyleInfo = style
    worksheet.add_table(table)

    # -------------------------
    # Formats
    # -------------------------
    for row in range(2, total_row):

        if isinstance(worksheet[f"E{row}"].value, datetime):
            worksheet[f"E{row}"].number_format = "yyyy-mm-dd"

        if isinstance(worksheet[f"F{row}"].value, datetime):
            worksheet[f"F{row}"].number_format = "yyyy-mm-dd"

        worksheet[f"I{row}"].number_format = "R #,##0.00"
        worksheet[f"J{row}"].number_format = "R #,##0.00"
        worksheet[f"K{row}"].number_format = "R #,##0.00"

    # -------------------------
    # Conditional Formatting
    # -------------------------
    green_fill = PatternFill(
        fill_type="solid",
        start_color="C6EFCE",
        end_color="C6EFCE",
    )

    yellow_fill = PatternFill(
        fill_type="solid",
        start_color="FFEB9C",
        end_color="FFEB9C",
    )

    red_fill = PatternFill(
        fill_type="solid",
        start_color="FFC7CE",
        end_color="FFC7CE",
    )

    # TD less CD
    worksheet.conditional_formatting.add(
        f"G2:G{total_row-1}",
        CellIsRule(
            operator="between",
            formula=["0", "14"],
            fill=green_fill,
        ),
    )

    worksheet.conditional_formatting.add(
        f"G2:G{total_row-1}",
        CellIsRule(
            operator="between",
            formula=["15", "30"],
            fill=yellow_fill,
        ),
    )

    worksheet.conditional_formatting.add(
        f"G2:G{total_row-1}",
        CellIsRule(
            operator="greaterThan",
            formula=["30"],
            fill=red_fill,
        ),
    )

    # EID less CD
    worksheet.conditional_formatting.add(
        f"H2:H{total_row-1}",
        CellIsRule(
            operator="lessThan",
            formula=["0"],
            fill=red_fill,
        ),
    )

    worksheet.conditional_formatting.add(
        f"H2:H{total_row-1}",
        CellIsRule(
            operator="between",
            formula=["0", "7"],
            fill=yellow_fill,
        ),
    )

    worksheet.conditional_formatting.add(
        f"H2:H{total_row-1}",
        CellIsRule(
            operator="greaterThan",
            formula=["7"],
            fill=green_fill,
        ),
    )

    # Balance Due
    worksheet.conditional_formatting.add(
        f"K2:K{total_row-1}",
        CellIsRule(
            operator="greaterThan",
            formula=["0"],
            fill=red_fill,
        ),
    )

    # -------------------------
    # Legend
    # -------------------------
    legend_column = "M"
    text_column = "N"

    # Title
    worksheet[f"{legend_column}1"] = "LEGEND"
    worksheet[f"{legend_column}1"].font = Font(bold=True, color="FFFFFF")
    worksheet[f"{legend_column}1"].fill = PatternFill(
        fill_type="solid",
        start_color="0F4C81",
        end_color="0F4C81",
    )
    worksheet[f"{legend_column}1"].alignment = Alignment(horizontal="center")

    worksheet.merge_cells("M1:N1")

    row = 3

    # -------------------------
    # TD less CD
    # -------------------------
    worksheet[f"N{row}"] = "TD less CD"
    worksheet[f"N{row}"].font = Font(bold=True)
    row += 1

    worksheet[f"M{row}"].fill = green_fill
    worksheet[f"N{row}"] = "0–14 days"
    row += 1

    worksheet[f"M{row}"].fill = yellow_fill
    worksheet[f"N{row}"] = "15–30 days"
    row += 1

    worksheet[f"M{row}"].fill = red_fill
    worksheet[f"N{row}"] = "Over 30 days"

    row += 2

    # -------------------------
    # EID less CD
    # -------------------------
    worksheet[f"N{row}"] = "EID less CD"
    worksheet[f"N{row}"].font = Font(bold=True)
    row += 1

    worksheet[f"M{row}"].fill = red_fill
    worksheet[f"N{row}"] = "Negative (Overdue)"
    row += 1

    worksheet[f"M{row}"].fill = yellow_fill
    worksheet[f"N{row}"] = "0–7 days"
    row += 1

    worksheet[f"M{row}"].fill = green_fill
    worksheet[f"N{row}"] = "More than 7 days"

    row += 2

    # -------------------------
    # Balance Due
    # -------------------------
    worksheet[f"N{row}"] = "Balance Due"
    worksheet[f"N{row}"].font = Font(bold=True)
    row += 1

    worksheet[f"M{row}"].fill = red_fill
    worksheet[f"N{row}"] = "Greater than R0.00"

    # Format legend
    worksheet.column_dimensions["M"].width = 6
    worksheet.column_dimensions["N"].width = 28

    for r in range(2, row + 1):
        worksheet[f"M{r}"].alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
    # -------------------------
    # Freeze header
    # -------------------------
    worksheet.freeze_panes = "A2"

    # -------------------------
    # Auto-size columns
    # -------------------------
    for column in worksheet.columns:
        max_length = 0
        letter = get_column_letter(column[0].column)

        for cell in column:
            try:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass

        worksheet.column_dimensions[letter].width = max_length + 3

    # Restore legend widths after auto-sizing
    worksheet.column_dimensions["M"].width = 10
    worksheet.column_dimensions["N"].width = 28

    # -------------------------
    # Save workbook
    # -------------------------
    output_folder = Path(REPORT_OUTPUT_FOLDER)
    output_folder.mkdir(parents=True, exist_ok=True)

    filename = REPORT_FILENAME.format(
        date=datetime.today().strftime("%Y-%m-%d")
    )

    file_path = output_folder / filename

    workbook.save(file_path)

    return file_path